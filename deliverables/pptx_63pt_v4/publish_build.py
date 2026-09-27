"""Restore hash-verified V4 report content, rebuild editable PPTX, and verify export.
Only document content is rebuilt. This does not rerun the numerical study.
"""
from pathlib import Path
import base64,copy,hashlib,json,math,os,subprocess,sys
from pptx import Presentation
import fitz

ROOT=Path('deliverables/pptx_63pt_v4')
BOOT=ROOT/'bootstrap'
M=json.loads((BOOT/'manifest.json').read_text())
def sha(b):return hashlib.sha256(b).hexdigest()
packed=base64.b64decode(''.join((BOOT/f'part{i}.b64').read_text().strip() for i in range(1,M['parts']+1)),validate=True)
assert sha(packed)==M['packed_sha256'],'Source archive hash mismatch'
import lzma
files=json.loads(lzma.decompress(packed))
assert set(files)==set(M['files'])
for name,content in files.items():
    assert isinstance(content,str) and sha(content.encode())==M['files'][name],name
    p=(ROOT/name).resolve()
    assert p.is_relative_to(ROOT.resolve()),'Unsafe path'
    p.parent.mkdir(parents=True,exist_ok=True);p.write_text(content,encoding='utf-8')
basebytes=Path(M['base_file']).read_bytes()
assert sha(basebytes)==M['base_sha256'],'Frozen V2 base mismatch'
base=json.loads(basebytes)
def restore(v):
    if isinstance(v,dict):
        if set(v)=={'__base__'}:
            out=base
            for k in v['__base__']:out=out[k]
            return copy.deepcopy(out)
        return {k:restore(x) for k,x in v.items()}
    if isinstance(v,list):return [restore(x) for x in v]
    return v
spec=restore(json.loads((ROOT/'revision.json').read_text()))
assert sha(json.dumps(spec,ensure_ascii=False,sort_keys=True).encode())==M['spec_canonical_sha256']
(ROOT/'deck_v4.json').write_text(json.dumps(spec,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
name='SSBB_V4_Submission_BaekSeungwoo'
pptx=ROOT/(name+'.pptx')
subprocess.run([sys.executable,str(ROOT/'build_pptx_v4.py'),'--spec',str(ROOT/'deck_v4.json'),'--out',str(pptx)],check=True)
r=Presentation(pptx)
assert len(r.slides)==len(spec['slides'])==58 and spec['main_count']==35
chart_count=value_count=cell_count=link_count=0
for sl,ds in zip(r.slides,spec['slides']):
    expected=[e for e in ds.get('elements',[]) if e['kind']=='chart']
    actual=[s.chart for s in sl.shapes if s.has_chart]
    assert len(expected)==len(actual)
    for ec,ac in zip(expected,actual):
        assert len(ec['series'])==len(ac.series)
        for es,ass in zip(ec['series'],ac.series):
            assert len(es['values'])==len(ass.values)
            assert all(math.isclose(float(a),float(b),rel_tol=1e-12,abs_tol=1e-12) for a,b in zip(es['values'],ass.values))
            value_count+=len(es['values'])
        chart_count+=1
    et=[e for e in ds.get('elements',[]) if e['kind']=='table'];at=[s.table for s in sl.shapes if s.has_table]
    assert len(et)==len(at)
    for e,t in zip(et,at):
        for i,row in enumerate(e['rows']):
            for j,v in enumerate(row):assert t.cell(i,j).text==str(v);cell_count+=1
    assert sl.has_notes_slide and len(sl.notes_slide.notes_text_frame.text)>100
    for s in sl.shapes:
        assert s.left>=-1000 and s.top>=-1000 and s.left+s.width<=r.slide_width+1000 and s.top+s.height<=r.slide_height+1000
    link_count+=len(sl._element.xpath('.//a:hlinkClick'))
assert (chart_count,value_count,cell_count,link_count)==(14,1035,1161,422)
subprocess.run(['libreoffice','-env:UserInstallation=file:///tmp/ssbb-v4-pdf','--headless','--convert-to','pdf','--outdir',str(ROOT),str(pptx)],check=True)
pdf=ROOT/(name+'.pdf')
assert pdf.exists() and pdf.stat().st_size>100000
with fitz.open(pdf) as doc:
    assert len(doc)==58
    assert all(len(p.get_text())>100 for p in doc)
    assert not any('\ufffd' in p.get_text() for p in doc)
qa={'status':'PASS','scope':'Remote document reconstruction and editable-object verification, not new numerical/model/HW validation. Full data re-audit is recorded in the local delivery package.','slides':58,'main_slides':35,'appendix_slides':23,'charts':chart_count,'chart_values_checked':value_count,'table_cells_checked':cell_count,'internal_links':link_count,'new_performance_simulations':False,'spec_canonical_sha256':M['spec_canonical_sha256']}
(ROOT/'remote_document_qa.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2)+'\n')
info={'source_commit':os.environ.get('GITHUB_SHA'),'workflow_run_url':os.environ.get('RUN_URL'),'document_version':'V4','dataset':'SSBB-NUM-20260923-v2 (63 baseline / 60 additional conditions)','publication_scope':'PPTX, PDF, glossary, review log, frozen document specification, builder and QA. Full frozen CSV/code and local audit package are also supplied as the conversation ZIP. No original corporate slides or font files are included.','artifacts':{p.name:{'bytes':p.stat().st_size,'sha256':sha(p.read_bytes())} for p in [pptx,pdf,ROOT/'deck_v4.json',ROOT/'SSBB_V4_Glossary_and_Reading_Guide.md',ROOT/'SSBB_V4_Review_and_Changes.md']}}
(ROOT/'publication.json').write_text(json.dumps(info,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(qa,ensure_ascii=False,indent=2))
