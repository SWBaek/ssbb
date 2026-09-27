"""Structural and numeric QA for the frozen presentation (not a model validation)."""
from pathlib import Path
from zipfile import ZipFile
import argparse, hashlib, json, math
from pptx import Presentation

def validate(pptx,spec):
    p=Path(pptx);d=json.loads(Path(spec).read_text(encoding='utf-8'));r=Presentation(p)
    assert len(r.slides)==38==len(d['slides'])
    counts={'slides':38,'main_slides':28,'appendix_slides':10,'native_charts':0,'native_tables':0,'notes':0,'verified_chart_points':0}
    for idx,(sl,s) in enumerate(zip(r.slides,d['slides']),1):
        visible='\n'.join(sh.text for sh in sl.shapes if sh.has_text_frame)
        assert s['title'] in visible, f'Title {idx}'
        for sh in sl.shapes:
            assert sh.left>=-200 and sh.top>=-200,(idx,sh.name,'negative bounds')
            assert sh.left+sh.width <= r.slide_width+4000,(idx,sh.name,'width')
            assert sh.top+sh.height <= r.slide_height+4000,(idx,sh.name,'height')
            if sh.has_table:counts['native_tables']+=1
        actual=[sh.chart for sh in sl.shapes if sh.has_chart];expected=[e for e in s.get('elements',[]) if e['kind']=='chart']
        assert len(actual)==len(expected)
        for ch,e in zip(actual,expected):
            counts['native_charts']+=1
            assert len(ch.series)==len(e['series'])
            for se,v in zip(ch.series,e['series']):
                assert se.name==v['name']
                assert len(se.values)==len(v['values'])
                for a,b in zip(se.values,v['values']):
                    assert math.isclose(a,b,rel_tol=1e-12,abs_tol=1e-12),(idx,a,b)
                    counts['verified_chart_points']+=1
        assert sl.has_notes_slide
        nt=sl.notes_slide.notes_text_frame.text
        assert '63-point main / 60-point holdout' in nt
        counts['notes']+=1
    assert counts['native_charts']==13
    with ZipFile(p) as z:
        assert z.testzip() is None
        assert not any(n.lower().endswith(('.ttf','.otf','.fntdata')) for n in z.namelist())
        counts['editable_chart_workbooks']=len([n for n in z.namelist() if n.startswith('ppt/embeddings/') and n.endswith('.xlsx')])
        assert counts['editable_chart_workbooks']==13
    return dict(status='passed',**counts,pptx_sha256=hashlib.sha256(p.read_bytes()).hexdigest(),spec_sha256=hashlib.sha256(Path(spec).read_bytes()).hexdigest(),scope='Presentation structure, source-derived numbers and notes; not product/BB approval.')
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--pptx',required=True);ap.add_argument('--spec',required=True);ap.add_argument('--out',required=True);a=ap.parse_args()
    result=validate(a.pptx,a.spec);Path(a.out).write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');print(json.dumps(result,ensure_ascii=False,indent=2))
