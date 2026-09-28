"""Validate editable objects and exported PDF against the generated V7 deck spec."""
from pathlib import Path
import json,math,hashlib,argparse
from pptx import Presentation
import fitz
from model import ROOT,assert_spec,SPEC

def verify(base):
    assert_spec();base=Path(base);ds=json.loads(base.with_suffix('.json').read_text());p=base.with_suffix('.pptx');prs=Presentation(p)
    assert len(prs.slides)==len(ds['slides'])
    nc=nv=nt=cell=links=0
    for sl,sd in zip(prs.slides,ds['slides']):
        charts=[x.chart for x in sl.shapes if x.has_chart];ce=[e for e in sd.get('elements',[]) if e['kind']=='chart'];assert len(charts)==len(ce)
        for c,e in zip(charts,ce):
            assert len(c.series)==len(e['series'])
            for s,z in zip(c.series,e['series']):
                assert len(s.values)==len(z['values'])
                assert all(math.isclose(float(a),float(b),rel_tol=1e-12,abs_tol=1e-12) for a,b in zip(s.values,z['values']))
                nv+=len(s.values)
            nc+=1
        tables=[x.table for x in sl.shapes if x.has_table];te=[e for e in sd.get('elements',[]) if e['kind']=='table'];assert len(tables)==len(te)
        for t,e in zip(tables,te):
            for i,row in enumerate(e['rows']):
                for j,x in enumerate(row):assert t.cell(i,j).text==str(x);cell+=1
            nt+=1
        assert sl.has_notes_slide and len(sl.notes_slide.notes_text_frame.text)>100
        for sh in sl.shapes:
            assert sh.left>=-1000 and sh.top>=-1000 and sh.left+sh.width<=prs.slide_width+1000 and sh.top+sh.height<=prs.slide_height+1000
        links+=len(sl._element.xpath('.//a:hlinkClick'))
    ids=[s['sid'] for s in ds['slides']];assert len(ids)==len(set(ids))
    for field,sid in [('test_map_page','tests'),('glossary_page','glossary'),('results_page','results')]:assert ids[ds[field]-1]==sid
    outside=[];glyph=[];inside_overflows=[]
    with fitz.open(base.with_suffix('.pdf')) as doc:
        assert len(doc)==len(ds['slides'])
        for i,page in enumerate(doc):
            assert len(page.get_text())>100
            if '\ufffd' in page.get_text():glyph.append(i+1)
            for b in page.get_text('dict')['blocks']:
                if 'lines' not in b:continue
                for ln in b['lines']:
                    for sp in ln['spans']:
                        x0,y0,x1,y1=sp['bbox']
                        if x0<-1 or y0<-1 or x1>page.rect.width+1 or y1>page.rect.height+1:outside.append(dict(page=i+1,text=sp['text'],bbox=sp['bbox']))
        assert not outside and not glyph,(outside,glyph)
    s=json.loads((ROOT/'results/summary.json').read_text());assert s['specification']==SPEC and s['campaign_runs']==502
    def h(p):return hashlib.sha256(p.read_bytes()).hexdigest()
    out=dict(status='PASS',scope='Editable-object/data binding and PDF page bounds; not a substitute for visual review or hardware validation',slides=len(ds['slides']),main_slides=ds['main_count'],appendix_slides=len(ds['slides'])-ds['main_count'],native_charts=nc,chart_values_checked=nv,native_tables=nt,table_cells_checked=cell,notes=len(ds['slides']),internal_links=links,pdf_text_outside_page=outside,replacement_glyph_pages=glyph,pptx_sha256=h(p),pdf_sha256=h(base.with_suffix('.pdf')),dataset=ds['dataset'])
    base.with_name(base.name+'_qa.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(out,ensure_ascii=False,indent=2));return out
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--base',default=str(ROOT/'report/SSBB_V7_Technical_Public'));a=ap.parse_args();verify(a.base)
