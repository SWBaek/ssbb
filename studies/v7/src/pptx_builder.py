"""Build an editable Six Sigma presentation from a frozen, traceable deck specification.
No simulation, threshold tuning, or result modification is performed here.
Usage: python build_pptx_v7.py --spec deck_v7.json --out SSBB_V7_Submission_BaekSeungwoo.pptx
"""
from __future__ import annotations
import argparse, json, math, hashlib
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_TICK_MARK
from pptx.chart.data import CategoryChartData
from pptx.oxml.xmlchemy import OxmlElement

C={'ink':'000000','muted':'555555','red':'C00000','teal':'008000','blue':'0056A6',
   'gray':'757575','light':'EFEFEF','line':'B8B8B8','white':'FFFFFF','amber':'805000','pale':'F3F3F3'}
FONT='Malgun Gothic'

def rgb(s): return RGBColor.from_string(C.get(s,s).replace('#',''))
def fill(shape,color):
    shape.fill.solid(); shape.fill.fore_color.rgb=rgb(color)
def font_run(r,size=13,bold=False,color='ink',family=FONT):
    r.font.name=family;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=rgb(color)
    pr=r._r.get_or_add_rPr()
    for n in ['a:ea','a:cs']:
        old=pr.find(n,pr.nsmap)
        if old is None: old=OxmlElement(n);pr.append(old)
        old.set('typeface',family)

def text(slide,txt,x,y,w,h,size=13,bold=False,color='ink',align='left',valign='top',margin=.025):
    sh=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=sh.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin);tf.margin_top=tf.margin_bottom=Inches(margin)
    tf.vertical_anchor={'top':MSO_ANCHOR.TOP,'middle':MSO_ANCHOR.MIDDLE,'bottom':MSO_ANCHOR.BOTTOM}[valign]
    for i,line in enumerate(str(txt).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
        p.space_before=Pt(0);p.space_after=Pt(2);p.line_spacing=1.08
        r=p.add_run();r.text=line;font_run(r,size,bold,color)
    return sh

def box(slide,x,y,w,h,color='light',line=None,radius=False):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    fill(shape,color)
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    for st in shape._element.xpath('./p:style'):shape._element.remove(st)
    if line:shape.line.color.rgb=rgb(line);shape.line.width=Pt(.65)
    else:shape.line.fill.background()
    if radius:
        try:shape.adjustments[0]=.08
        except Exception:pass
    return shape

def line(slide,x,y,x2,y2,color='line',width=1,arrow=False,dash=False):
    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(x2),Inches(y2))
    sh.line.color.rgb=rgb(color);sh.line.width=Pt(width)
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    for st in sh._element.xpath('./p:style'):sh._element.remove(st)
    if arrow:
        e=OxmlElement('a:tailEnd');e.set('type','triangle');e.set('w','sm');e.set('len','sm');sh.line._get_or_add_ln().append(e)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    return sh

def table(slide,e):
    rows=e['rows'];x,y,w,h=e['x'],e['y'],e['w'],e['h'];nr=len(rows);nc=len(rows[0])
    sh=slide.shapes.add_table(nr,nc,Inches(x),Inches(y),Inches(w),Inches(h));t=sh.table
    widths=e.get('widths',[1]*nc);sw=sum(widths)
    for j,v in enumerate(widths):t.columns[j].width=Inches(w*v/sw)
    heights=e.get('heights',[1]*nr);hs=sum(heights)
    for i,v in enumerate(heights):t.rows[i].height=Inches(h*v/hs)
    for i,row in enumerate(rows):
        for j,v in enumerate(row):
            cell=t.cell(i,j);cell.text='';cell.margin_left=cell.margin_right=Inches(.07);cell.margin_top=cell.margin_bottom=Inches(.045)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            color='ink';bg='white'
            if i==0:bg='light'
            if i in e.get('highlight_rows',[]):bg='pale';color=e.get('highlight_color','teal')
            style=e.get('cell_styles',{}).get(f'{i},{j}',{})
            bg=style.get('fill',bg);color=style.get('color',color)
            cell.fill.solid();cell.fill.fore_color.rgb=rgb(bg)
            for k,s in enumerate(str(v).split('\n')):
                p=cell.text_frame.paragraphs[0] if k==0 else cell.text_frame.add_paragraph()
                p.space_before=Pt(0);p.space_after=Pt(1);p.line_spacing=1.05
                p.alignment=PP_ALIGN.CENTER if i==0 or j in e.get('center_cols',[]) else PP_ALIGN.LEFT
                r=p.add_run();r.text=s;font_run(r,e.get('size',12),i==0 or j in e.get('bold_cols',[]),color)
            tcpr=cell._tc.get_or_add_tcPr()
            for name in ['a:lnL','a:lnR','a:lnT','a:lnB']:
                ln=OxmlElement(name);ln.set('w','6350');sf=OxmlElement('a:solidFill');cr=OxmlElement('a:srgbClr');cr.set('val',C['line']);sf.append(cr);ln.append(sf);tcpr.append(ln)
    return sh

def chart(slide,e):
    data=CategoryChartData()
    cats=e['categories'];skip=e.get('cat_skip',1)
    data.categories=[v if skip<=1 or i%skip==0 or i==len(cats)-1 else '' for i,v in enumerate(cats)]
    for s in e['series']:data.add_series(s['name'],s['values'])
    typ={'line':XL_CHART_TYPE.LINE,'column':XL_CHART_TYPE.COLUMN_CLUSTERED,'bar':XL_CHART_TYPE.BAR_CLUSTERED}[e.get('type','line')]
    ch=slide.shapes.add_chart(typ,Inches(e['x']),Inches(e['y']),Inches(e['w']),Inches(e['h']),data).chart
    ch.has_title=False;ch.has_legend=e.get('legend',len(e['series'])>1)
    if ch.has_legend:
        ch.legend.position=XL_LEGEND_POSITION.BOTTOM;ch.legend.include_in_layout=False
        ch.legend.font.name=FONT;ch.legend.font.size=Pt(e.get('legend_size',10))
    ch.font.name=FONT;ch.font.size=Pt(e.get('size',10))
    ca=ch.category_axis;va=ch.value_axis
    for ax in [ca,va]:
        ax.tick_labels.font.name=FONT;ax.tick_labels.font.size=Pt(e.get('size',10));ax.tick_labels.font.color.rgb=rgb('muted')
        ax.major_tick_mark=XL_TICK_MARK.NONE;ax.minor_tick_mark=XL_TICK_MARK.NONE
        ax.format.line.color.rgb=rgb('line')
    va.minimum_scale=e.get('min',0)
    if 'max' in e:va.maximum_scale=e['max']
    if 'major' in e:va.major_unit=e['major']
    va.has_major_gridlines=True;va.major_gridlines.format.line.color.rgb=rgb('line');va.major_gridlines.format.line.width=Pt(.45)
    va.tick_labels.number_format=e.get('format','0.0');va.tick_labels.number_format_is_linked=False
    ca.has_major_gridlines=False
    from pptx.enum.chart import XL_TICK_LABEL_POSITION
    ca.tick_label_position=XL_TICK_LABEL_POSITION.LOW
    if e.get('cat_skip'):
        for n in ['c:tickLblSkip','c:tickMarkSkip']:
            z=OxmlElement(n);z.set('val','1')
            end=ca._element.find('c:noMultiLvlLbl', ca._element.nsmap)
            if end is not None:ca._element.insert(ca._element.index(end), z)
            else:ca._element.append(z)
    for idx,s in enumerate(ch.series):
        col=e['series'][idx].get('color',['gray','teal','red','blue'][idx%4])
        s.format.line.color.rgb=rgb(col);s.format.line.width=Pt(e['series'][idx].get('width',1.8))
        if e.get('discrete') and not e['series'][idx].get('dash'): s.format.line.fill.background()
        s.format.fill.solid();s.format.fill.fore_color.rgb=rgb(col)
        if e.get('type','line')=='line':
            try:
                from pptx.enum.chart import XL_MARKER_STYLE
                s.marker.style=XL_MARKER_STYLE.NONE
                if e.get('markers') or (e.get('discrete') and not e['series'][idx].get('dash')):
                    s.marker.style=XL_MARKER_STYLE.CIRCLE;s.marker.size=4;s.marker.format.fill.solid();s.marker.format.fill.fore_color.rgb=rgb(col);s.marker.format.line.color.rgb=rgb(col)
            except Exception:pass
        if e['series'][idx].get('dash'):
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            s.format.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    plot=ch.plots[0]
    if e.get('labels'):
        plot.has_data_labels=True;dl=plot.data_labels;dl.font.name=FONT;dl.font.size=Pt(10);dl.number_format=e.get('label_format','0.00');dl.position=XL_LABEL_POSITION.OUTSIDE_END
    if e.get('type') in ['bar','column']:plot.gap_width=70
    return ch

def elements(slide,els):
    for e in els:
        k=e['kind']
        if k=='text':text(slide,e['text'],e['x'],e['y'],e['w'],e['h'],e.get('size',13),e.get('bold',False),e.get('color','ink'),e.get('align','left'),e.get('valign','top'))
        elif k=='box':
            box(slide,e['x'],e['y'],e['w'],e['h'],'white' if e.get('fill','light')=='light' else e.get('fill','white'),e.get('line','line'),False)
            if e.get('text'):text(slide,e['text'],e['x']+.09,e['y']+.07,e['w']-.18,e['h']-.14,e.get('size',13),e.get('bold',False),e.get('color','ink'),e.get('align','left'),e.get('valign','middle'))
        elif k=='line':line(slide,e['x'],e['y'],e['x2'],e['y2'],e.get('color','line'),e.get('width',1),e.get('arrow',False),e.get('dash',False))
        elif k=='table':table(slide,e)
        elif k=='chart':chart(slide,e)
        elif k=='metric':
            box(slide,e['x'],e['y'],e['w'],e['h'],'white','line')
            text(slide,e['label'],e['x']+.16,e['y']+.12,e['w']-.32,.35,11,False,'muted')
            text(slide,e['value'],e['x']+.16,e['y']+.57,e['w']-.32,.6,e.get('value_size',25),True,e.get('color','red'))
            if e.get('detail'):text(slide,e['detail'],e['x']+.16,e['y']+1.16,e['w']-.32,e['h']-1.2,10.5,False,'muted')
        else:raise ValueError(k)

def build(spec,out):
    global FONT
    FONT=spec.get('font',FONT);r=Presentation();r.slide_width=Inches(spec.get('width',10.833333));r.slide_height=Inches(7.5)
    cp=r.core_properties;cp.title=spec['title'];cp.subject='Six Sigma BB / Report V7, verified terminology and fresh V7 switching data';cp.author=spec['author'];cp.keywords='DMAIC,THD,V2G,numerical simulation';cp.comments='Source-grounded presentation. Not hardware validation or BB approval.'
    total=len(spec['slides']);main=spec.get('main_count',28)
    for i,s in enumerate(spec['slides'],1):
        sl=r.slides.add_slide(r.slide_layouts[6]);fill(sl.background,'white') if False else None
        if s.get('cover'):
            text(sl,'Six Sigma BB 과제 보고서',.53,.4,9.5,.45,20,True)
            line(sl,.0,1.03,10.833333,1.03,'ink',.7)
            text(sl,s['title'],.62,1.62,9.58,1.42,27,True)
            text(sl,s.get('subtitle','수치해석 기반 고조파 전류제어 검토'),.65,3.35,9.4,.7,16,True)
            box(sl,.65,4.08,9.55,.48,'light','line')
            text(sl,'Define  →  Measure  →  Analyze  →  Improve  →  Control',.8,4.17,9.2,.32,14,True,align='center')
            text(sl,s['author_line'],.68,5.25,9.2,.52,15)
            text(sl,s['date_line'],.68,6.02,9.2,.65,11,color='muted')
            line(sl,.4,7.02,10.44,7.02,'ink',.45)
            text(sl,'문서·수치 데이터 V7  |  240 Vac · 48 Arms | PWM·제어 62.5 kHz / 16 μs | Peak dq  |  실제 제품 검증과 구분',.43,7.12,9.8,.22,8.2,color='muted')
        else:
            phase=s.get('stage','A')
            text(sl,s['title'],.42,.20,8.15,.46,20,True)
            for j,k in enumerate('DMAIC'):
                xx=8.72+j*.325
                sh=sl.shapes.add_shape(MSO_SHAPE.CHEVRON,Inches(xx),Inches(.22),Inches(.38),Inches(.32))
                fill(sh,'D6D6D6' if k==phase else 'F4F4F4');sh.line.color.rgb=rgb('line');sh.line.width=Pt(.45)
                try:sh.adjustments[0]=.28
                except Exception:pass
                text(sl,k,xx+.035,.221,.26,.30,10.5,k==phase,'ink' if k==phase else 'muted','center','middle',.005)
            line(sl,0,.79,10.833333,.79,'ink',.65)
            if s.get('subtitle'):
                box(sl,.40,.95,10.05,.34,'light','line')
                text(sl,'■ '+s['subtitle'],.50,.98,9.86,.28,12.3,True)
            if s.get('message'):
                text(sl,s['message'],.48,1.43,9.89,.63,13.2)
            elements(sl,s.get('elements',[]))
            line(sl,.40,7.03,10.45,7.03,'ink',.45)
            text(sl,s.get('source',''),.43,7.09,9.4,.28,7.7,color='muted')
            text(sl,f'{i}/{total}',9.84,7.09,.57,.23,8.5,False,'ink','right')
        notes=s.get('notes','')
        sl.notes_slide.notes_text_frame.text=(f"[{i}/{total}] {s['title']}\n"+notes+'\n\n근거: '+s.get('source','')+'\n데이터 출처: '+spec['dataset']+'\n주의: 수치모델 연구이며 실제 제품 인증/BB 승인을 의미하지 않습니다.')
    for i,(sl,sd) in enumerate(zip(r.slides,spec['slides'])):
        if sd.get('cover'): continue
        targets=[('시험 구성',spec.get('test_map_page',3)),('용어집',spec.get('glossary_page',36))]
        if i+1>=spec.get('glossary_page',36): targets.append(('본문 결과',spec.get('results_page',25)))
        x=7.92 if len(targets)==2 else 6.65
        for label,page in targets:
            sh=text(sl,f'{label} {page}쪽',x,7.33,1.24,.14,7.8,False,'blue','right',margin=0)
            sh.click_action.target_slide=r.slides[page-1]
            x+=1.26
        # Keep phase chevrons clickable without making page labels look like data.
        for sh in sl.shapes:
            if sh.has_text_frame and sh.text in 'DMAIC' and len(sh.text)==1:
                page=spec.get('stage_pages',{}).get(sh.text)
                if page: sh.click_action.target_slide=r.slides[page-1]
    out=Path(out);out.parent.mkdir(parents=True,exist_ok=True);r.save(out)
    audit={'slides':total,'main_slides':main,'appendix_slides':total-main,'native_tables':sum(sum(sh.has_table for sh in s.shapes) for s in r.slides),'native_charts':sum(len(s.shapes._spTree.xpath('.//c:chart')) for s in r.slides),'dataset':spec['dataset'],'pptx_sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'spec_sha256':hashlib.sha256(json.dumps(spec,ensure_ascii=False,sort_keys=True).encode()).hexdigest()}
    out.with_suffix('.build.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(audit,ensure_ascii=False,indent=2))
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--spec',default='deck_v7.json');ap.add_argument('--out',default='SSBB_V7_Submission_BaekSeungwoo.pptx');a=ap.parse_args();build(json.loads(Path(a.spec).read_text()),a.out)
